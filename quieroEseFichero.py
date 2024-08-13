import os
import requests
from bs4 import BeautifulSoup
import re
from urllib.parse import urljoin, urlparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
import xlrd

# Función para aplicar colores en terminal
def color_text(text, color):
    colors = {
        'red': '\033[91m',
        'green': '\033[92m',
        'yellow': '\033[93m',
        'blue': '\033[94m',
        'magenta': '\033[95m',
        'cyan': '\033[96m',
        'white': '\033[97m',
        'reset': '\033[0m'
    }
    return f"{colors[color]}{text}{colors['reset']}"

# Variables globales para almacenar los resultados
found_files = []
visited = set()

def download_file(url, output_folder):
    local_filename = url.split('/')[-1]
    local_path = os.path.join(output_folder, local_filename)
    
    print(color_text(f"Descargando archivo desde: {url}", 'blue'))
    
    try:
        with requests.get(url, stream=True) as r:
            r.raise_for_status()
            with open(local_path, 'wb') as f:
                for chunk in r.iter_content(chunk_size=8192):
                    f.write(chunk)
        print(color_text(f"Archivo descargado: {local_path}", 'green'))
        return local_path
    except Exception as e:
        print(color_text(f"Error al descargar {url}: {e}", 'red'))
        return None

def extract_metadata(file_path):
    metadata = {}
    try:
        if file_path.endswith('.pdf'):
            with open(file_path, 'rb') as f:
                pdf = PdfReader(f)
                info = pdf.metadata
                if info:
                    metadata['Title'] = getattr(info, 'title', None)
                    metadata['Author'] = getattr(info, 'author', None)
                    metadata['Creator'] = getattr(info, 'creator', None)
                    metadata['Producer'] = getattr(info, 'producer', None)
                    metadata['Subject'] = getattr(info, 'subject', None)
                    metadata['Keywords'] = getattr(info, 'keywords', None)
                else:
                    print(color_text(f"No se encontraron metadatos en el archivo PDF: {file_path}", 'yellow'))
        elif file_path.endswith('.docx'):
            doc = Document(file_path)
            core_props = doc.core_properties
            metadata['Title'] = core_props.title
            metadata['Author'] = core_props.author
            metadata['Last Modified By'] = core_props.last_modified_by
            metadata['Created'] = core_props.created
            metadata['Modified'] = core_props.modified
            metadata['Subject'] = core_props.subject
            metadata['Keywords'] = core_props.keywords
        elif file_path.endswith('.xlsx'):
            wb = load_workbook(file_path, read_only=True)
            core_props = wb.properties
            metadata['Title'] = core_props.title
            metadata['Author'] = core_props.creator
            metadata['Last Modified By'] = core_props.lastModifiedBy
            metadata['Created'] = core_props.created
            metadata['Modified'] = core_props.modified
            metadata['Keywords'] = core_props.keywords
        elif file_path.endswith('.xls'):
            wb = xlrd.open_workbook(file_path)
            metadata['Title'] = wb.props.title
            metadata['Author'] = wb.props.author
            metadata['Created'] = wb.props.created
            metadata['Modified'] = wb.props.modified
    except Exception as e:
        print(color_text(f"Error al extraer metadatos de {file_path}: {e}", 'red'))
    return metadata

def find_documents(url, domain, output_folder):
    global found_files, visited

    if url in visited:
        return
    visited.add(url)
    
    print(color_text(f"Accediendo a: {url}", 'blue'))

    try:
        response = requests.get(url)
        response.raise_for_status()
    except requests.RequestException as e:
        print(color_text(f"Error al acceder a {url}: {e}", 'red'))
        return

    soup = BeautifulSoup(response.content, 'html.parser')

    # Encuentra todos los enlaces en la página
    links = soup.find_all('a', href=True)

    # Expresión regular para documentos .doc, .docx, .pdf, .xls, .xlsx
    doc_pattern = re.compile(r'.*\.(doc|docx|pdf|xls|xlsx)$', re.IGNORECASE)

    for link in links:
        href = link['href']
        full_url = urljoin(url, href)
        parsed_url = urlparse(full_url)

        # Solo sigue enlaces que pertenecen al dominio principal y subdominios
        if domain in parsed_url.netloc:
            if doc_pattern.match(parsed_url.path):
                print(color_text(f"Documento encontrado: {full_url}", 'green'))
                found_files.append((full_url, None))  # Se almacena la URL para descargar más tarde
            else:
                if full_url not in visited:
                    find_documents(full_url, domain, output_folder)
        else:
            print(color_text(f"Ignorado enlace fuera del dominio: {full_url}", 'yellow'))

def download_and_analyze_file(output_folder):
    display_files()
    if not found_files:
        return
    try:
        choice = int(input(color_text("\nIntroduce el número del archivo que quieres descargar: ", 'cyan')))
        if 1 <= choice <= len(found_files):
            url, _ = found_files[choice - 1]
            local_path = download_file(url, output_folder)
            found_files[choice - 1] = (url, local_path)  # Actualiza la ruta local

            # Extraer metadatos después de la descarga
            if local_path:
                metadata = extract_metadata(local_path)
                if metadata:
                    print(color_text("\nMetadatos extraídos:", 'magenta'))
                    for key, value in metadata.items():
                        print(color_text(f"{key}: {value}", 'green'))
                else:
                    print(color_text("No se encontraron metadatos o no se pudieron extraer.", 'yellow'))
        else:
            print(color_text("Opción no válida.", 'red'))
    except ValueError:
        print(color_text("Por favor, introduce un número válido.", 'red'))

def display_files():
    if not found_files:
        print(color_text("No se encontraron archivos.", 'yellow'))
    else:
        print(color_text("\nLista de archivos encontrados:", 'cyan'))
        for idx, (url, path) in enumerate(found_files, start=1):
            print(color_text(f"{idx}. {url}", 'green'))

def parallel_find_documents(urls, domain, output_folder):
    with ThreadPoolExecutor(max_workers=5) as executor:
        futures = [executor.submit(find_documents, url, domain, output_folder) for url in urls]
        for future in as_completed(futures):
            future.result()

def confirm_exit():
    confirm = input(color_text("\n¿Estás seguro de que deseas salir? [Y/N]: ", 'red')).upper()
    return confirm == 'Y'

def generate_report():
    report_path = "reporte_ficheros.txt"
    try:
        with open(report_path, 'w') as report_file:
            report_file.write("Reporte de Ficheros Encontrados y Metadatos\n")
            report_file.write("========================================\n\n")
            if found_files:
                for url, path in found_files:
                    if path:
                        report_file.write(f"URL: {url}\n")
                        report_file.write(f"Local: {path}\n")
                        metadata = extract_metadata(path)
                        if metadata:
                            report_file.write("Metadatos:\n")
                            for key, value in metadata.items():
                                report_file.write(f"  {key}: {value}\n")
                        else:
                            report_file.write("No se encontraron metadatos.\n")
                        report_file.write("\n")
            else:
                report_file.write("No se encontraron archivos para generar el reporte.\n")
        print(color_text(f"Reporte generado en: {report_path}", 'cyan'))
    except Exception as e:
        print(color_text(f"Error al generar el reporte: {e}", 'red'))

def main_menu():
    while True:
        print(color_text("\n+----------------------------------+", 'cyan'))
        print(color_text("|         QUIERO ESE FICHERO       |", 'yellow'))
        print(color_text("|       Author: Oscar Martínez     |", 'yellow'))
        print(color_text("+----------------------------------+\n", 'cyan'))
        
        print(color_text("E) xplore Website for Documents", 'green'))
        print(color_text("L) ist Found Files", 'green'))
        print(color_text("D) ownload and Analyze a Selected File", 'green'))
        print(color_text("R) eport Generation", 'green'))
        print(color_text("Q) uit\n", 'red'))

        command = input(color_text("Command : ", 'magenta')).upper()

        if command == 'E':
            base_url = input(color_text("\nIntroduce la URL del sitio web a analizar: ", 'cyan'))
            output_folder = 'documentos_descargados'
            if not os.path.exists(output_folder):
                os.makedirs(output_folder)
            domain = urlparse(base_url).netloc
            parallel_find_documents([base_url], domain, output_folder)
        elif command == 'L':
            display_files()
        elif command == 'D':
            output_folder = 'documentos_descargados'
            if not os.path.exists(output_folder):
                os.makedirs(output_folder)
            download_and_analyze_file(output_folder)
        elif command == 'R':
            generate_report()
        elif command == 'Q':
            if confirm_exit():
                print(color_text("Saliendo del programa. ¡Hasta luego!", 'red'))
                break
        else:
            print(color_text("Opción no válida. Por favor, selecciona una opción del menú.", 'red'))

if __name__ == "__main__":
    main_menu()
